"""Excel writer built on openpyxl, with light formatting for readability."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

import pandas as pd


def write(df: pd.DataFrame, path: str | Path | io.BytesIO, sheet_name: str = "Data") -> Any:
    """Write ``df`` to an .xlsx file (or an in-memory buffer) and return it."""
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:  # pragma: no cover - dependency is pinned
        raise ImportError(
            "Writing Excel files needs openpyxl. Install it with: pip install openpyxl"
        ) from exc

    # A BytesIO target has no parent directory to create.
    if isinstance(path, (str, Path)):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name[:31])
        _format_sheet(writer.sheets[sheet_name[:31]], df)
    return path


def _format_sheet(worksheet, df: pd.DataFrame) -> None:
    """Bold the header, freeze it, and widen columns to fit their content."""
    from openpyxl.styles import Alignment, Font, PatternFill

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="374151")
    for cell in worksheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    worksheet.freeze_panes = "A2"

    for index, column in enumerate(df.columns, start=1):
        longest = max(
            [len(str(column))]
            + [len(str(value)) for value in df[column].head(500).fillna("")]
        )
        worksheet.column_dimensions[
            worksheet.cell(row=1, column=index).column_letter
        ].width = min(max(longest + 2, 10), 60)